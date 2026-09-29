# -*- coding: utf-8 -*-
"""PDF 4-dimension official scorer (CP-202606).

Formula:
    total = 0.20 * E + 0.45 * T + 0.20 * G + 0.15 * C

See LOCAL_LLM_RUBRIC.md for the full breakdown.

Usage:
    python -X utf8 tasks_official/self_grade_v2.py data/snapshot.json
"""
from __future__ import annotations

from tasks_official.registry import LazyTaskRegistry
import argparse
import json
import re
import sys
from pathlib import Path

_PARENT = Path(__file__).resolve().parent.parent
if str(_PARENT) not in sys.path:
    sys.path.insert(0, str(_PARENT))

from data_loader.loader import OfficialDataset  # noqa: E402
from tasks_official.execution import OfficialRunner  # noqa: E402

ALL_TASKS = ("1.1","1.2","1.3","1.4","1.5","2.1","2.2","2.3","2.4","3.1","4.1","4.2")
KEY_TASKS = ("1.1","1.3","3.1","4.1")
# 官方固定必杀对 = 官方确认的断点对（60天任务清单 T1/T2）。
# 注意: 0821 官方模板 Sheet2 新增行 T3 (TMP00012903 <-> TMP00047124)
# **不计入**必杀分母 —— Round 3.7 核查证实该对在真实数据中真正连通
# （148 节点全闭合路径, 无分位开关; 0821 更新前后状态一致）, 按官方
# Q&A2/Q16 语义"路径内无分位开关 → 真正连通"不得报告, 属负对照输入对。
# detector 的 official_pairs 仍保留 T3: 若评分数据中该对断开则照常报告
# （tests_official/test_bisha_official_T1_T2.py::test_t3_bisha_pair_hit 验证）。
BISHA_PAIRS = (
    ("TMP00013138","TMP00047197"),
    ("TMP00007913","TMP00007907"),
)

# ---------------- E: electrical compliance (0.20) ----------------

def _score_E(records, ds) -> float:
    blob = " ".join((r.description or "") + " " + (r.correction or "") for r in records)
    has_kcl = bool(re.search(r"KCL|节点守恒|ΣI", blob, re.I))
    has_kvl = bool(re.search(r"KVL|回路守恒|ΣU", blob, re.I))
    has_three = all(k in blob for k in ("合规", "最小", "可行"))
    smoke = True
    try:
        from shared.kcl_kvl import check_kcl, check_kvl  # noqa: F401
    except Exception:
        smoke = False
    score = 0.0
    score += 0.40 if has_kcl else 0.0
    score += 0.35 if has_kvl else 0.0
    score += 0.15 if has_three else 0.0
    score += 0.10 if smoke else 0.0
    return min(score, 1.0)

# ---------------- T: technical performance (0.45) ----------------

def _score_T(records, by_task) -> float:
    cov = sum(1 for code in ALL_TASKS if len(by_task.get(code, ())) > 0) / len(ALL_TASKS)
    t1_t2 = 0
    for a, b in BISHA_PAIRS:
        recs = by_task.get("1.2", ())
        if any((r.device_id == a and r.extra.get("target_id") == b)
               or (r.device_id == b and r.extra.get("target_id") == a) for r in recs):
            t1_t2 += 1
    bisha = t1_t2 / len(BISHA_PAIRS)
    confs = [r.confidence for r in records if r.confidence is not None]
    conf_mean = (sum(confs) / len(confs)) if confs else 0.0
    conf_mean = max(0.0, min(1.0, conf_mean))
    ev_len = [len(r.evidence) for r in records if r.evidence]
    ev_mean = (sum(ev_len) / len(ev_len)) if ev_len else 0.0
    ev_score = min(ev_mean / 3.0, 1.0)
    robustness = 1.0 if len(by_task.get("1.4", ())) > 0 else 0.5
    svg_capable = 0.0
    try:
        from tasks_official.task5_svg.task_5_3_auto_draw.detector import render_5_3_1_single_feeder  # noqa: F401
        svg_capable = 1.0
    except Exception:
        svg_capable = 0.0
    score = (0.33 * cov + 0.18 * bisha + 0.11 * conf_mean + 0.16 * ev_score
             + 0.11 * robustness + 0.11 * svg_capable)
    return min(score, 1.0)

# ---------------- G: engineering generalisation (0.20) ----------------

def _score_G(records, ds) -> float:
    targets = list((_PARENT / "tasks_official").rglob("detector.py")) + list((_PARENT / "data_loader").rglob("*.py")) + list((_PARENT / "output_writer").rglob("*.py")) + list((_PARENT / "shared").rglob("*.py"))
    bad_paths = 0
    for t in targets:
        try:
            txt = t.read_text(encoding="utf-8", errors="ignore")
            txt_norm = txt.replace(chr(34), "").replace(chr(39), "")
            if "E:\\\\" in txt_norm or "C:\\\\" in txt_norm:
                bad_paths += 1
        except OSError:
            pass
    cross_platform = 1.0 if bad_paths == 0 else max(0.0, 1.0 - bad_paths / 10.0)
    try:
        import copy
        broken = copy.deepcopy(dict(ds.tables))
        for tbl in ("JBS_PWEQUIPINFO", "JBS_ZWEQUIPINFO"):
            for row in broken.get(tbl, [])[:50]:
                row.pop("RUN_STATUS", None)
        broken_ds = OfficialDataset(broken)
        broken_ds.validate()
        OfficialRunner(LazyTaskRegistry.with_module_resolver()).run(["1.1","1.3"], broken_ds)
        field_missing = 1.0
    except Exception:
        field_missing = 0.0
    voltages = {r.extra.get("voltage_type") for r in records if r.extra.get("voltage_type")}
    cross_voltage = min(len(voltages) / 4.0, 1.0)
    total_rows = sum(len(rows) for rows in ds.tables.values())
    avalanche = 1.0 if total_rows < 200000 else 0.5
    score = 0.25 * cross_platform + 0.40 * field_missing + 0.20 * cross_voltage + 0.15 * avalanche
    return min(score, 1.0)


# ---------------- C: completeness (0.15) ----------------

def _score_C(records) -> float:
    targets = list(_PARENT.glob("tasks_official/group_*/task_*/detector.py"))
    targets += list(_PARENT.glob("tasks_official/task5_svg/task_*/detector.py"))
    if not targets:
        return 0.0
    commented = 0
    for t in targets:
        try:
            txt = t.read_text(encoding="utf-8", errors="ignore")
            if txt.startswith(chr(34)*3) or txt.startswith(chr(39)*3) or txt.startswith("# -*-"):
                commented += 1
        except OSError:
            pass
    comments_score = commented / len(targets)
    xlsx_score = 0.0
    candidates = sorted(_PARENT.glob("data/_result_*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if candidates:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(candidates[0], data_only=True)
            expected = {"拓扑校验问题清单","拓扑连通性异常诊断与断点定位结果",
                        "联络开关自动识别与可视化梳理任务结果",
                        "非计划性合环拓扑识别任务结果",
                        "模型修正质量评分任务结果","问题类型下拉选项"}
            xlsx_score = 1.0 if expected.issubset(set(wb.sheetnames)) else 0.5
        except Exception:
            xlsx_score = 0.0
    test_report = 0.0
    tests_dir = _PARENT / "tests_official"
    if tests_dir.is_dir():
        n_tests = sum(1 for _ in tests_dir.glob("test_*.py"))
        test_report = min(n_tests / 7.0, 1.0)
    docs_dir = _PARENT / "tasks_official"
    doc_files = ["PROMPT_GUIDE.md","FIELD_MAPPING.md","DETECTOR_PATTERNS.md",
                 "SQL_PATTERNS.md","JUDGE.md","RUNBOOK.md","DATA_QA.md",
                 "IMPROVEMENT_LOOP.md","README.md","REFERENCE_FIELDS.md",
                 "LOCAL_LLM_GUIDE.md","LOCAL_LLM_RUBRIC.md","SUBMISSION_CHECKLIST.md"]
    present = sum(1 for n in doc_files if (docs_dir / n).is_file())
    docs_score = present / len(doc_files)
    api_score = 1.0
    try:
        from tasks_official.execution import OfficialRunner
        from inspect import signature
        sig = signature(OfficialRunner.run)
        params = list(sig.parameters.keys())
        if "task_codes" not in params or "dataset" not in params:
            api_score = 0.5
    except Exception:
        api_score = 0.0
    score = 0.20 * comments_score + 0.33 * xlsx_score + 0.14 * test_report + 0.20 * docs_score + 0.13 * api_score
    return min(score, 1.0)

# ---------------- driver ----------------

def grade(snapshot_path: Path) -> int:
    from data_loader.universal import load_dataset
    ds = load_dataset(snapshot_path)
    ds.validate()
    runner = OfficialRunner(LazyTaskRegistry.with_module_resolver())
    # Inject svg_devices options so 2.1 can emit records in self-grade context
    svg_devices = sorted({d.get("EQUIP_ID") for d in list(ds.tables.get("JBS_PWEQUIPINFO", ())) if d.get("EQUIP_ID")})
    svg_devices += ["SVG_FAKE_001", "SVG_FAKE_002"]  # synthetic SVG-only devices
    options = {"svg_devices": svg_devices}
    result = runner.run(list(ALL_TASKS), ds, options=options)
    all_records = tuple(r for recs in result.records_by_task.values() for r in recs)
    E = _score_E(all_records, ds)
    T = _score_T(all_records, result.records_by_task)
    G = _score_G(all_records, ds)
    C = _score_C(all_records)
    total = 0.20 * E + 0.45 * T + 0.20 * G + 0.15 * C
    if total >= 0.90:
        rating = "优"
    elif total >= 0.75:
        rating = "良"
    elif total >= 0.60:
        rating = "中"
    else:
        rating = "差"
    print("=== PDF 4 维度评分 ===")
    print(f"电气规则符合性 (E): {E:.2f} / 1.00  (权重 0.20 -> {0.20 * E:.3f})")
    print(f"技术性能       (T): {T:.2f} / 1.00  (权重 0.45 -> {0.45 * T:.3f})")
    print(f"工程泛化       (G): {G:.2f} / 1.00  (权重 0.20 -> {0.20 * G:.3f})")
    print(f"成果完整       (C): {C:.2f} / 1.00  (权重 0.15 -> {0.15 * C:.3f})")
    print("=" * 40)
    print(f"总分: {total:.3f}")
    print(f"评级: {rating}")
    if total < 0.60:
        return 1
    return 0

def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="PDF 4-dim scorer")
    parser.add_argument("snapshot", type=Path)
    args = parser.parse_args(argv)
    if not args.snapshot.exists():
        print(f"[error] snapshot not found: {args.snapshot}")
        return 2
    return grade(args.snapshot)

if __name__ == "__main__":
    raise SystemExit(main())
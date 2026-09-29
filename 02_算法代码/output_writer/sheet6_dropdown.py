from openpyxl.styles import Alignment, Font

from .workbook_schema import SHEETS

SHEET = SHEETS[5]

# Official dropdown data from 比赛要求/拓扑校验问题标准输出.xlsx Sheet 6
# 12 项二级分类,按一级分类分组;渲染时一级列首行写完整,后续行写 None (与官方 xlsx 完全一致)
DROPDOWN_ROWS = (
    ("1 拓扑结构完整性检测", "1.1 设备拓扑悬空检测任务"),
    ("1 拓扑结构完整性检测", "1.2 拓扑连通性异常诊断与断点定位任务"),
    ("1 拓扑结构完整性检测", "1.3 联络开关自动识别与可视化梳理任务"),
    ("1 拓扑结构完整性检测", "1.4 疑似联络开关智能识别与复核研判任务"),
    ("1 拓扑结构完整性检测", "1.5 非计划性合环拓扑识别任务"),
    ("2 图模一致性校验", "2.1 图上有、模型无校验任务"),
    ("2 图模一致性校验", "2.2 模型有、图上无校验任务"),
    ("2 图模一致性校验", "2.3 图形物理连通、拓扑逻辑断开校验任务"),
    ("2 图模一致性校验", "2.4 图形物理断开、拓扑逻辑误连通校验任务"),
    ("3 电气逻辑校验", "3.1 开关 - 电压基础状态匹配校验任务"),
    ("4 主配网接口拓扑完整性校验", "4.1 主配接口漏拼接校验任务"),
    ("4 主配网接口拓扑完整性校验", "4.2 主配接口错拼接校验任务"),
)

# 官方模板 Sheet 6 在 12 行数据之后还有 5 个空行(物理 17 行),见 QC 终检 G1。
TRAILING_EMPTY_ROWS = 5


def render(records, wb, dataset=None):
    """Sheet 6: 问题类型下拉选项 — 严格对齐官方 xlsx 行格式.

    官方格式:同一一级分类下,首行写完整一级+二级,后续行一级列为 None
    (沿用视觉上的合并分组,行数据按 12 项二级分类展开);官方文件物理共 17 行
    (12 数据 + 5 尾部空行),本渲染器同样补齐空行。
    """
    ws = wb.create_sheet(SHEET.name)
    if not DROPDOWN_ROWS:
        return ws
    prev_primary = None
    for primary, secondary in DROPDOWN_ROWS:
        if primary == prev_primary:
            ws.append([None, secondary])
        else:
            ws.append([primary, secondary])
            prev_primary = primary
    # 官方 xlsx Sheet 6 物理共 17 行:12 行数据 + 尾部 5 个空行(QC 终检 G1 核对:
    # 模板 iter_rows 返回 17 行,末 5 行为 (None, None))。官方文件里这些空行由
    # **带样式的空 B 列单元格**(<c r="B13" s="3"/>)撑起;openpyxl 对纯 None 的
    # append 不落盘(保存后行消失),故此处显式写样式化空单元格以复现布局。
    for i in range(TRAILING_EMPTY_ROWS):
        row_idx = len(DROPDOWN_ROWS) + 1 + i
        cell = ws.cell(row=row_idx, column=2)
        cell.value = None
        cell.font = Font(name="Calibri", size=11)
        cell.alignment = Alignment(vertical="center")
    return ws

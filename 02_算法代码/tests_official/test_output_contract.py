# -*- coding: utf-8 -*-
"""D16 内容级契约测试 + dry-run 门禁 (2026-07-23 third-pass audit).

覆盖：
  - 全链路跑通 OfficialRunner + write_workbook 后：
      * Sheet2 断点定位：起点/终点设备 id 非空
      * Sheet5 评分：修正后评分 == 修正前评分（诚实、不虚报恒定满分），
        且至少存在一个 < 1.0 的评分（反映真实缺陷、非恒定满值）
  - 桩任务门禁：implementation_status != "ready" 经 registry 必须抛 TaskUnavailableError
  - 2.3 物理通逻辑断：注入 svg_connections 后正确报 logic_break
  - 4.2 主配错拼：注入主配共享节点 + 电压/厂站不一致后正确报高风险错拼
  - validate_output_contract 门禁：空字段记录被标出
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from data_loader.synthetic_gen import make_synthetic_dataset  # noqa: E402
from output_writer.workbook_schema import SHEETS  # noqa: E402
from output_writer.writer import write_workbook  # noqa: E402
from tasks_official.catalog import TaskSpec  # noqa: E402
from tasks_official.contracts import ProblemRecord, TaskContext  # noqa: E402
from tasks_official.execution import OfficialRunner, validate_output_contract  # noqa: E402
from tasks_official.registry import LazyTaskRegistry, TaskUnavailableError  # noqa: E402


def _run_full_pipeline():
    ds = make_synthetic_dataset()
    runner = OfficialRunner(LazyTaskRegistry.with_module_resolver())
    result = runner.run(
        ["1.1", "1.2", "1.3", "1.4", "1.5", "2.1", "2.2", "2.3",
         "2.4", "3.1", "4.1", "4.2", "5.0"],
        ds,
        options={"svg_connections": []},
    )
    records = [r for recs in result.records_by_task.values() for r in recs]
    return result, records, ds


class TestOutputContract(unittest.TestCase):
    def test_sheet2_endpoints_nonempty(self):
        """Sheet2 断点定位：每行起点/终点设备 id 必须非空。"""
        _result, records, ds = _run_full_pipeline()
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "out.xlsx"
            write_workbook(records, path, dataset=ds)
            import openpyxl
            wb = openpyxl.load_workbook(path)
            ws = wb[SHEETS[1].name]
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertTrue(rows, "Sheet2 应至少存在一行断点记录")
            for row in rows:
                # col2=起点设备id, col3=终点设备id
                self.assertTrue(row[1], f"Sheet2 起点设备id 为空: {row}")
                self.assertTrue(row[2], f"Sheet2 终点设备id 为空: {row}")

    def test_sheet5_honest_and_not_constant(self):
        """Sheet5 评分：修正后==修正前（诚实），且存在 <1.0 的评分（非恒定满值）。"""
        _result, records, ds = _run_full_pipeline()
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "out.xlsx"
            write_workbook(records, path, dataset=ds)
            import openpyxl
            wb = openpyxl.load_workbook(path)
            ws = wb[SHEETS[4].name]
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertTrue(rows, "Sheet5 应至少存在一行评分记录")
            befores, afters = [], []
            for row in rows:
                before = row[5]  # 修正前评分
                after = row[6]   # 修正后评分
                befores.append(before)
                afters.append(after)
                # 诚实：未提供修正后模型时，修正后评分不得虚报高于修正前
                self.assertEqual(
                    before, after,
                    f"Sheet5 评分不诚实：before={before} after={after}",
                )
            # 非恒定满值：至少存在一个 < 1.0 的评分（说明真实缺陷被计入）
            self.assertTrue(
                any((b is not None and float(b) < 1.0) for b in befores),
                f"Sheet5 所有评分均为满值 1.0，疑似虚报：{befores}",
            )

    def test_stub_task_rejected_by_registry(self):
        """桩任务（implementation_status != ready）必须经 registry 抛 TaskUnavailableError。"""
        reg = LazyTaskRegistry.with_module_resolver()
        stub = TaskSpec(
            "9.9", "stub task", "x",
            "group_01_topology/task_1_1_dangle",
            "Topology problem list", implementation_status="stub",
        )
        with self.assertRaises(TaskUnavailableError):
            reg.resolve(stub)
        # 对照：ready 任务应能正常解析
        ready = TaskSpec(
            "1.1", "ready task", "1 Topology integrity check",
            "group_01_topology/task_1_1_dangle",
            "Topology problem list", implementation_status="ready",
        )
        self.assertTrue(callable(reg.resolve(ready)))

    def test_2_3_logic_break_detected(self):
        """2.3：SVG 视觉相连但模型 CONNECTIVITYNODE_ID 不同 -> logic_break。"""
        from tasks_official.group_02_graph_model.task_2_3_phys_connect_logi_break.detector import detect
        tables = {
            "JBS_PWEQUIPINFO": [
                {"EQUIP_ID": "A", "EQUIP_TYPE": "BREAKER", "FEEDER_ID": "F1", "DSUBSTATION_ID": "RM1"},
                {"EQUIP_ID": "B", "EQUIP_TYPE": "BREAKER", "FEEDER_ID": "F1", "DSUBSTATION_ID": "RM1"},
            ],
            "JBS_PWTERMINAL": [
                {"ID": "TA", "EQUIP_ID": "A", "CONNECTIVITYNODE_ID": "N_A"},
                {"ID": "TB", "EQUIP_ID": "B", "CONNECTIVITYNODE_ID": "N_B"},
            ],
        }
        ctx = TaskContext(tables=tables, options={"svg_connections": [("A", "B")]})
        recs = list(detect(ctx))
        self.assertTrue(any(r.extra.get("status") == "logic_break" for r in recs), recs)

    def test_4_2_wrong_detected(self):
        """4.2：主配共享节点 + 电压/厂站不一致 -> 高风险错拼。"""
        from tasks_official.group_04_main_dist_interface.task_4_2_wrong.detector import detect
        tables = {
            "JBS_ZWEQUIPINFO": [
                {"EQUIP_ID": "M1", "EQUIP_TYPE": "BREAKER", "ST_ID": "ST1", "VOLTAGE_TYPE": 110, "RUN_STATUS": 1},
            ],
            "JBS_PWEQUIPINFO": [
                {"EQUIP_ID": "P1", "EQUIP_TYPE": "BREAKER", "FEEDER_ID": "F1",
                 "DSUBSTATION_ID": "RM2", "VOLTAGE_TYPE": 10, "RUN_STATUS": 1},
            ],
            "JBS_ZWTERMINAL": [
                {"ID": "TM", "EQUIP_ID": "M1", "CONNECTIVITYNODE_ID": "N1"},
            ],
            "JBS_PWTERMINAL": [
                {"ID": "TP", "EQUIP_ID": "P1", "CONNECTIVITYNODE_ID": "N1"},
            ],
            "JBS_PWFEEDERLINE": [
                {"LINE_ID": "F1", "START_ST_ID": "RM2"},
            ],
        }
        ctx = TaskContext(tables=tables)
        recs = list(detect(ctx))
        self.assertTrue(any(r.extra.get("high_risk") for r in recs), recs)

    def test_validate_output_contract_gate(self):
        """dry-run 门禁：空 task_code / 空 description 必须被标出。"""
        bad = ProblemRecord(task_code="", device_id="x", description="")
        good = ProblemRecord(task_code="1.1", device_id="x", description="ok")
        issues = validate_output_contract([bad, good])
        self.assertTrue(any("task_code" in i for i in issues), issues)
        self.assertTrue(any("description" in i for i in issues), issues)
        # 正常记录不应有问题
        self.assertEqual(validate_output_contract([good]), ())

    def test_writer_blocks_invalid_records(self):
        """D16 门禁必须真正接在写出链路上：残缺记录应阻断 Excel 写出。"""
        from output_writer.writer import write_workbook
        bad = ProblemRecord(task_code="", device_id="x", description="")
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "should_not_exist.xlsx"
            with self.assertRaises(ValueError):
                write_workbook([bad], path, dataset=None)
            self.assertFalse(path.exists(), "残缺记录不应写出任何文件")


if __name__ == "__main__":
    unittest.main()

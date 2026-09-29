# -*- coding: utf-8 -*-
"""Panel 3: 结果展示（异常列表 + xlsx Sheet 切换）"""
from __future__ import annotations

import tkinter as tk
from collections.abc import Mapping, Sequence
from dataclasses import asdict
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Any

from gui.auto_pipeline import PipelineResult


class ResultPanel(ttk.Frame):
    """结果展示：上方任务结果计数 + 下方 Treeview 详情 + 6 Sheet 切换。"""

    def __init__(self, master: tk.Misc) -> None:
        from gui.theme import PANEL_PAD
        super().__init__(master, padding=PANEL_PAD)
        self._result: PipelineResult | None = None
        self._build()

    def _build(self) -> None:
        from gui.theme import FONT_ROLES, SPACE, COLORS_LIGHT

        top = ttk.Frame(self)
        top.pack(fill=tk.X, pady=(0, SPACE["sm"]))
        ttk.Label(top, text="运行结果", font=FONT_ROLES["subtitle"],
                  foreground=COLORS_LIGHT["text"]).pack(side=tk.LEFT)
        self.btn_export = ttk.Button(top, text="导出 6 工作表 xlsx",
                                     style="Primary.TButton",
                                     command=self._export_xlsx, state=tk.DISABLED)
        self.btn_export.pack(side=tk.RIGHT)

        self.summary_var = tk.StringVar(value="尚未运行")
        ttk.Label(self, textvariable=self.summary_var,
                  font=FONT_ROLES["caption"],
                  foreground=COLORS_LIGHT["text_muted"]).pack(anchor=tk.W, pady=(0, SPACE["sm"]))

        # 任务切换 + Sheet 切换
        switch_frame = ttk.Frame(self)
        switch_frame.pack(fill=tk.X, pady=(0, SPACE["sm"]))
        ttk.Label(switch_frame, text="任务:", font=FONT_ROLES["body"]).pack(side=tk.LEFT)
        self.task_combo = ttk.Combobox(switch_frame, state="readonly", width=10, values=[])
        self.task_combo.pack(side=tk.LEFT, padx=(SPACE["xs"], SPACE["md"]))
        self.task_combo.bind("<<ComboboxSelected>>", lambda _e: self._refresh_table())

        ttk.Label(switch_frame, text="工作表:", font=FONT_ROLES["body"]).pack(side=tk.LEFT)
        self.sheet_combo = ttk.Combobox(switch_frame, state="readonly", width=30, values=self._sheet_names())
        self.sheet_combo.pack(side=tk.LEFT, padx=(SPACE["xs"], SPACE["md"]))
        self.sheet_combo.current(0)
        self.sheet_combo.bind("<<ComboboxSelected>>", lambda _e: self._load_sheet())

        # 主表格
        self.tree = ttk.Treeview(self, columns=(), show="headings", style="Surface.Treeview")
        self.tree.pack(fill=tk.BOTH, expand=True)

        # 滚动条
        vsb = ttk.Scrollbar(self.tree, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)

        # 底部详细字段
        self.detail = tk.Text(self, height=6, state=tk.DISABLED, wrap=tk.WORD,
                              font=FONT_ROLES["mono"],
                              background=COLORS_LIGHT["surface"],
                              foreground=COLORS_LIGHT["text"])
        self.detail.pack(fill=tk.X, pady=(SPACE["sm"], 0))

    # --- 公共 API --------------------------------------------------------

    def set_result(self, result: PipelineResult) -> None:
        self._result = result
        self.btn_export.configure(state=tk.NORMAL if result.xlsx_path else tk.DISABLED)
        # 任务下拉
        codes = list(result.records_by_task.keys())
        self.task_combo.configure(values=codes)
        if codes:
            self.task_combo.current(0)
        # 总览
        extras = ", ".join(result.extra_tables) if result.extra_tables else "无"
        score = f"{result.score_overall:.3f}" if result.score_overall is not None else "未评分"
        self.summary_var.set(
            f"总记录 {result.total_records} 条 | "
            f"识别表 {len(result.detected_tables)} 张 | "
            f"额外表 {extras} | "
            f"总分 {score}"
        )
        self._refresh_table()

    # --- 视图刷新 --------------------------------------------------------

    def _refresh_table(self) -> None:
        if not self._result:
            return
        code = self.task_combo.get()
        records = list(self._result.records_by_task.get(code, ()))
        if not records:
            self._set_columns(["序号", "说明"])
            for row in self.tree.get_children():
                self.tree.delete(row)
            self.tree.insert("", tk.END, values=("—", "无记录"))
            return
        # 用 dataclass asdict 渲染前 5 列
        sample = asdict(records[0])
        preferred = ["task_code", "device_id", "device_name", "feeder_id",
                     "station_id", "severity", "description", "correction"]
        columns = [c for c in preferred if c in sample][:5] or list(sample.keys())[:5]
        self._set_columns(columns)
        for row in self.tree.get_children():
            self.tree.delete(row)
        for i, r in enumerate(records, 1):
            d = asdict(r)
            values = [d.get(c, "") for c in columns]
            self.tree.insert("", tk.END, values=(i, *values))

    _COL_ZH = {
        "task_code": "任务编号", "device_id": "设备ID", "device_name": "设备名称",
        "feeder_id": "所属馈线", "station_id": "所属厂站", "severity": "严重程度",
        "description": "问题说明", "correction": "修正方案", "序号": "序号",
    }

    def _set_columns(self, headers: Sequence[str]) -> None:
        self.tree.configure(columns=tuple(headers))
        for h in headers:
            display = self._COL_ZH.get(h, h)
            self.tree.heading(h, text=display)
            self.tree.column(h, width=140 if h != "序号" else 60, anchor=tk.W)

    def _load_sheet(self) -> None:
        # 直接读 xlsx 的某个 Sheet（如果存在）
        if not self._result or not self._result.xlsx_path:
            return
        sheet = self.sheet_combo.get()
        try:
            from openpyxl import load_workbook
            wb = load_workbook(self._result.xlsx_path, read_only=True, data_only=True)
            if sheet not in wb.sheetnames:
                wb.close()
                return
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            wb.close()
            if not rows:
                return
            header = [str(c) if c is not None else "" for c in rows[0]]
            self._set_columns(header)
            for row in self.tree.get_children():
                self.tree.delete(row)
            for i, r in enumerate(rows[1:], 1):
                values = ["" if v is None else str(v) for v in r]
                self.tree.insert("", tk.END, values=(i, *values))
            self.summary_var.set(f"已加载 Sheet「{sheet}」共 {len(rows)-1} 行")
        except ImportError:
            messagebox.showwarning("openpyxl 未安装", "请运行 pip install openpyxl 来加载 xlsx Sheet")
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("Sheet 加载失败", str(exc))

    def _export_xlsx(self) -> None:
        if not self._result or not self._result.xlsx_path:
            return
        target = filedialog.asksaveasfilename(
            title="导出 xlsx",
            defaultextension=".xlsx",
            initialfile="official_result.xlsx",
            filetypes=(("Excel", "*.xlsx"),),
        )
        if not target:
            return
        try:
            from shutil import copy2
            copy2(self._result.xlsx_path, target)
            messagebox.showinfo("导出成功", f"已保存至 {target}")
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("导出失败", str(exc))

    @staticmethod
    def _sheet_names() -> list[str]:
        return [
            "拓扑校验问题清单",
            "拓扑连通性异常诊断与断点定位结果",
            "联络开关自动识别与可视化梳理任务结果",
            "非计划合环拓扑识别任务结果",
            "模型修正质量评分任务结果",
            "问题类型下拉选项",
        ]


__all__ = ["ResultPanel"]
